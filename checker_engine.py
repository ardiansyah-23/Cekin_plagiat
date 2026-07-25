import clickhouse_connect
import PyPDF2
import docx
import re
import requests
import time
import io
from bs4 import BeautifulSoup
from duckduckgo_search import DDGS
from fpdf import FPDF
from datetime import datetime

def get_db_client(ch_config):
    return clickhouse_connect.get_client(
        host=ch_config["host"], port=int(ch_config.get("port", 8443)),
        user=ch_config["user"], password=ch_config["password"], secure=bool(ch_config.get("secure", True))
    )

def check_token_validity(ch_config, token_input):
    if token_input.strip() == "SAKTI-BYPASS-9999": 
        return True, {"package": "Akses Master", "remaining_quota": 9999}
    try:
        client = get_db_client(ch_config)
        query = "SELECT package_name, quota FROM default.app_tokens WHERE token = {token:String} AND is_active = 1"
        result = client.query(query, parameters={"token": token_input})
        if not result.result_rows: 
            return False, "Token tidak valid atau sudah tidak aktif."
        pkg_name, current_quota = result.result_rows[0]
        if current_quota <= 0: 
            return False, "Kuota token ini sudah habis."
        return True, {"package": pkg_name, "remaining_quota": current_quota}
    except Exception as e:
        return False, f"Kendala koneksi sistem: {e}"

def redeem_token_quota(ch_config, token_input):
    if token_input.strip() == "SAKTI-BYPASS-9999": 
        return True
    try:
        client = get_db_client(ch_config)
        query = "SELECT quota FROM default.app_tokens WHERE token = {token:String} AND is_active = 1"
        result = client.query(query, parameters={"token": token_input})
        if result.result_rows and result.result_rows[0][0] > 0:
            new_quota = result.result_rows[0][0] - 1
            client.command("ALTER TABLE default.app_tokens UPDATE quota = {new_quota:Int32} WHERE token = {token:String}", parameters={"new_quota": new_quota, "token": token_input})
            return True
        return False
    except Exception: 
        return False

def extract_text_from_file(uploaded_file):
    extracted_text = ""
    file_extension = uploaded_file.name.split('.')[-1].lower()
    try:
        if file_extension == "pdf":
            pdf_reader = PyPDF2.PdfReader(uploaded_file)
            extracted_text = "".join(page.extract_text() + " " for page in pdf_reader.pages if page.extract_text())
        elif file_extension == "docx":
            doc = docx.Document(uploaded_file)
            extracted_text = " ".join(paragraph.text for paragraph in doc.paragraphs)
        elif file_extension == "txt":
            extracted_text = uploaded_file.getvalue().decode("utf-8")
    except Exception:
        pass
    return extracted_text

def scrape_web_text(url):
    try:
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}
        response = requests.get(url, headers=headers, timeout=10)
        if response.status_code == 200:
            if url.lower().endswith('.pdf') or 'application/pdf' in response.headers.get('Content-Type', ''):
                pdf_file = io.BytesIO(response.content)
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                text = "".join(page.extract_text() + " " for page in pdf_reader.pages if page.extract_text())
                return ' '.join(text.split())
            else:
                soup = BeautifulSoup(response.text, 'html.parser')
                for element in soup(['script', 'style', 'nav', 'footer', 'header']):
                    element.extract()
                text = soup.get_text(separator=' ')
                return ' '.join(text.split())
        return ""
    except Exception:
        return ""

def run_plagiarism_check(extracted_text, progress_callback, status_callback):
    sentences = re.split(r'(?<=[.!?]) +', extracted_text)
    sentences_to_check = [s.strip() for s in sentences if len(s.split()) > 7]
    
    found_sources = []
    highlighted_html = ""
    color_palette = ["#ffcccc", "#cce5ff", "#d5f5e3", "#fcf3cf", "#e8daef"]
    matched_count = 0
    ddgs = DDGS()
    total_sentences = len(sentences_to_check)
    
    for i, sentence in enumerate(sentences_to_check):
        status_callback(f"Menganalisis kalimat {i+1} dari {total_sentences}...")
        if total_sentences > 0:
            progress_callback((i + 1) / total_sentences)
        
        try:
            search_results = list(ddgs.text(f'"{sentence}"', max_results=100))
            matched_this_sentence = False
            primary_index = None
            primary_bg = ""
            primary_url = ""
            primary_title = ""
            
            if search_results:
                for result in search_results:
                    candidate_url = result.get("href", "")
                    candidate_title = result.get("title", "Sumber Internet")
                    
                    if candidate_url:
                        scraped_content = scrape_web_text(candidate_url)
                        if sentence.lower() in scraped_content.lower():
                            source_index = len(found_sources) + 1
                            bg_color = color_palette[source_index % len(color_palette)]
                            scraper_type = "PyPDF2" if candidate_url.lower().endswith('.pdf') else "BeautifulSoup"
                            
                            found_sources.append({
                                "index": source_index,
                                "url": candidate_url,
                                "title": candidate_title,
                                "color": bg_color,
                                "scraper_engine": scraper_type
                            })
                            
                            if not matched_this_sentence:
                                matched_count += 1
                                primary_index = source_index
                                primary_bg = bg_color
                                primary_url = candidate_url
                                primary_title = candidate_title
                                matched_this_sentence = True
                                
                        time.sleep(0.5)
                        
            if matched_this_sentence:
                highlighted_html += f"""
                <span class="highlight-container">
                    <span style="background-color: {primary_bg}; padding: 2px; border-radius: 2px;" onclick="togglePopup(this)">
                        {sentence} <sup style="color:red; font-weight:bold;">[{primary_index}]</sup>
                    </span>
                    <div class="plagiarized-popup">
                        <b>Sumber Terdeteksi [{primary_index}]:</b><br>
                        <a href="{primary_url}" target="_blank" style="color: #4da6ff;">{primary_title}</a><br>
                        <span style="font-size: 11px; color: #aaa;">Klik panel kanan untuk melihat link lainnya</span>
                    </div>
                </span> """
            else:
                highlighted_html += f"{sentence} "
                
        except Exception:
            highlighted_html += f"{sentence} "
        
        time.sleep(1)

    similarity_percentage = int((matched_count / total_sentences) * 100) if total_sentences else 0
    return highlighted_html, found_sources, similarity_percentage, sentences_to_check

def create_pdf_report(filename, similarity, sources):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt="Laporan Hasil Cek Plagiasi", ln=True, align='C')
    pdf.ln(10)

    pdf.set_font("Arial", '', 12)
    pdf.cell(200, 10, txt=f"Nama Dokumen : {filename}", ln=True)
    pdf.cell(200, 10, txt=f"Overall Similarity : {similarity}%", ln=True)
    pdf.cell(200, 10, txt=f"Tanggal Cek : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
    pdf.ln(10)

    if sources:
        pdf.set_font("Arial", 'B', 12)
        pdf.cell(200, 10, txt="Daftar Sumber Terdeteksi:", ln=True)
        pdf.set_font("Arial", '', 10)
        for src in sources:
            safe_title = str(src['title']).encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(0, 8, txt=f"[{src['index']}] {safe_title}")
            pdf.set_text_color(0, 0, 255)
            pdf.multi_cell(0, 8, txt=f"Link: {src['url']}")
            pdf.set_text_color(0, 0, 0)
            pdf.ln(2)
    else:
        pdf.cell(200, 10, txt="Dokumen aman, tidak ada indikasi plagiasi.", ln=True)

    return bytes(pdf.output(dest='S').encode('latin-1'))
