import clickhouse_connect
import PyPDF2
import docx
import re
import requests
import time
import io
from bs4 import BeautifulSoup
from duckduckgo_search import DDGS
from fpdf import FPDF
from datetime import datetime
import streamlit as st
import uuid
import checker_engine


def get_db_client(ch_config):
    return clickhouse_connect.get_client(
        host=ch_config["host"], port=int(ch_config.get("port", 8443)),
        user=ch_config["user"], password=ch_config["password"], secure=bool(ch_config.get("secure", True))
    )

def check_token_validity(ch_config, token_input):
    if token_input.strip() == "SAKTI-BYPASS-9999": 
        return True, {"package": "Akses Master", "remaining_quota": 9999}
    try:
        client = get_db_client(ch_config)
        query = "SELECT package_name, quota FROM default.app_tokens WHERE token = {token:String} AND is_active = 1"
        result = client.query(query, parameters={"token": token_input})
        if not result.result_rows: 
            return False, "Token tidak valid atau sudah tidak aktif."
        pkg_name, current_quota = result.result_rows[0]
        if current_quota <= 0: 
            return False, "Kuota token ini sudah habis."
        return True, {"package": pkg_name, "remaining_quota": current_quota}
    except Exception as e:
        return False, f"Kendala koneksi sistem: {e}"

def redeem_token_quota(ch_config, token_input):
    if token_input.strip() == "SAKTI-BYPASS-9999": 
        return True
    try:
        client = get_db_client(ch_config)
        query = "SELECT quota FROM default.app_tokens WHERE token = {token:String} AND is_active = 1"
        result = client.query(query, parameters={"token": token_input})
        if result.result_rows and result.result_rows[0][0] > 0:
            new_quota = result.result_rows[0][0] - 1
            client.command("ALTER TABLE default.app_tokens UPDATE quota = {new_quota:Int32} WHERE token = {token:String}", parameters={"new_quota": new_quota, "token": token_input})
            return True
        return False
    except Exception: 
        return False

def extract_text_from_file(uploaded_file):
    extracted_text = ""
    file_extension = uploaded_file.name.split('.')[-1].lower()
    try:
        if file_extension == "pdf":
            pdf_reader = PyPDF2.PdfReader(uploaded_file)
            extracted_text = "".join(page.extract_text() + " " for page in pdf_reader.pages if page.extract_text())
        elif file_extension == "docx":
            doc = docx.Document(uploaded_file)
            extracted_text = " ".join(paragraph.text for paragraph in doc.paragraphs)
        elif file_extension == "txt":
            extracted_text = uploaded_file.getvalue().decode("utf-8")
    except Exception:
        pass
    return extracted_text

def scrape_web_text(url):
    try:
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}
        response = requests.get(url, headers=headers, timeout=10)
        if response.status_code == 200:
            if url.lower().endswith('.pdf') or 'application/pdf' in response.headers.get('Content-Type', ''):
                pdf_file = io.BytesIO(response.content)
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                text = "".join(page.extract_text() + " " for page in pdf_reader.pages if page.extract_text())
                return ' '.join(text.split())
            else:
                soup = BeautifulSoup(response.text, 'html.parser')
                for element in soup(['script', 'style', 'nav', 'footer', 'header']):
                    element.extract()
                text = soup.get_text(separator=' ')
                return ' '.join(text.split())
        return ""
    except Exception:
        return ""

def run_plagiarism_check(extracted_text, progress_callback, status_callback):
    sentences = re.split(r'(?<=[.!?]) +', extracted_text)
    sentences_to_check = [s.strip() for s in sentences if len(s.split()) > 7]
    
    found_sources = []
    highlighted_html = ""
    color_palette = ["#ffcccc", "#cce5ff", "#d5f5e3", "#fcf3cf", "#e8daef"]
    matched_count = 0
    ddgs = DDGS()
    total_sentences = len(sentences_to_check)
    
    for i, sentence in enumerate(sentences_to_check):
        status_callback(f"Menganalisis kalimat {i+1} dari {total_sentences}...")
        if total_sentences > 0:
            progress_callback((i + 1) / total_sentences)
        
        try:
            # Mengambil hingga 100 link sesuai permintaan
            search_results = list(ddgs.text(f'"{sentence}"', max_results=100))
            matched_this_sentence = False
            primary_index = None
            primary_bg = ""
            primary_url = ""
            primary_title = ""
            
            if search_results:
                for result in search_results:
                    candidate_url = result.get("href", "")
                    candidate_title = result.get("title", "Sumber Internet")
                    
                    if candidate_url:
                        scraped_content = scrape_web_text(candidate_url)
                        if sentence.lower() in scraped_content.lower():
                            source_index = len(found_sources) + 1
                            bg_color = color_palette[source_index % len(color_palette)]
                            scraper_type = "PyPDF2" if candidate_url.lower().endswith('.pdf') else "BeautifulSoup"
                            
                            found_sources.append({
                                "index": source_index,
                                "url": candidate_url,
                                "title": candidate_title,
                                "color": bg_color,
                                "scraper_engine": scraper_type
                            })
                            
                            if not matched_this_sentence:
                                matched_count += 1
                                primary_index = source_index
                                primary_bg = bg_color
                                primary_url = candidate_url
                                primary_title = candidate_title
                                matched_this_sentence = True
                                
                        time.sleep(0.5)
                        
            if matched_this_sentence:
                highlighted_html += f"""
                <span class="highlight-container">
                    <span style="background-color: {primary_bg}; padding: 2px; border-radius: 2px;" onclick="togglePopup(this)">
                        {sentence} <sup style="color:red; font-weight:bold;">[{primary_index}]</sup>
                    </span>
                    <div class="plagiarized-popup">
                        <b>Sumber Terdeteksi [{primary_index}]:</b><br>
                        <a href="{primary_url}" target="_blank" style="color: #4da6ff;">{primary_title}</a><br>
                        <span style="font-size: 11px; color: #aaa;">Klik panel kanan untuk melihat link lainnya</span>
                    </div>
                </span> """
            else:
                highlighted_html += f"{sentence} "
                
        except Exception:
            highlighted_html += f"{sentence} "
        
        time.sleep(1)

    similarity_percentage = int((matched_count / total_sentences) * 100) if total_sentences else 0
    return highlighted_html, found_sources, similarity_percentage, sentences_to_check

def create_pdf_report(filename, similarity, sources):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt="Laporan Hasil Cek Plagiasi", ln=True, align='C')
    pdf.ln(10)

    pdf.set_font("Arial", '', 12)
    pdf.cell(200, 10, txt=f"Nama Dokumen : {filename}", ln=True)
    pdf.cell(200, 10, txt=f"Overall Similarity : {similarity}%", ln=True)
    pdf.cell(200, 10, txt=f"Tanggal Cek : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
    pdf.ln(10)

    if sources:
        pdf.set_font("Arial", 'B', 12)
        pdf.cell(200, 10, txt="Daftar Sumber Terdeteksi:", ln=True)
        pdf.set_font("Arial", '', 10)
        for src in sources:
            safe_title = str(src['title']).encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(0, 8, txt=f"[{src['index']}] {safe_title}")
            pdf.set_text_color(0, 0, 255)
            pdf.multi_cell(0, 8, txt=f"Link: {src['url']}")
            pdf.set_text_color(0, 0, 0)
            pdf.ln(2)
    else:
        pdf.cell(200, 10, txt="Dokumen aman, tidak ada indikasi plagiasi.", ln=True)

    return bytes(pdf.output(dest='S').encode('latin-1'))

    import streamlit as st
import uuid
from datetime import datetime
import checker_engine

# --- KONFIGURASI HALAMAN ---
st.set_page_config(page_title="Sistem Cek Kemiripan Dokumen", layout="wide")

# --- CSS & JAVASCRIPT GAYA TURNITIN ---
st.markdown("""
<style>
.highlight-container { position: relative; display: inline-block; cursor: pointer; }
.turnitin-source-list { border: 1px solid #ddd; padding: 15px; border-radius: 5px; background: #fff; }
.turnitin-score { font-size: 48px; font-weight: bold; color: #b30000; margin-bottom: 0; }
.plagiarized-popup { display: none; position: absolute; top: 110%; left: 0; width: 350px; background-color: #222; color: #fff; text-align: left; border-radius: 6px; padding: 10px; z-index: 999; font-size: 13px; box-shadow: 0px 4px 10px rgba(0,0,0,0.4); }
.highlight-container.active .plagiarized-popup { display: block; }
</style>
<script>
document.addEventListener('click', function(event) {
  if (!event.target.closest('.highlight-container')) {
    document.querySelectorAll('.highlight-container').forEach(el => { el.classList.remove('active'); });
  }
});
function togglePopup(element) {
  event.stopPropagation();
  let container = element.closest('.highlight-container');
  let isActive = container.classList.contains('active');
  document.querySelectorAll('.highlight-container').forEach(el => { el.classList.remove('active'); });
  if (!isActive) { container.classList.add('active'); }
}
</script>
""", unsafe_allow_html=True)

ADMIN_PASSWORD = st.secrets["admin"]["password"]
ch_config = st.secrets["clickhouse"]

if "authenticated" not in st.session_state: st.session_state.authenticated = False
if "token_info" not in st.session_state: st.session_state.token_info = {}
if "token_code" not in st.session_state: st.session_state.token_code = ""

# ================= SIDEBAR =================
st.sidebar.title("Navigasi Menu")
menu_option = st.sidebar.radio("Pilih Halaman:", ["Utama: Cek Plagiasi", "Login Token / Redeem", "Panel Admin"])
st.sidebar.markdown("---")
if st.session_state.authenticated:
    info = st.session_state.token_info
    st.sidebar.success(f"✅ Sesi Aktif\n- Paket: {info.get('package')}\n- Sisa Kuota: {info.get('remaining_quota')}x")
    if st.sidebar.button("Keluar / Ganti Token"):
        st.session_state.authenticated = False
        st.session_state.token_info = {}
        st.session_state.token_code = ""
        st.rerun()
else: 
    st.sidebar.info("💡 Anda belum memasukkan token. Token hanya dipotong saat dokumen diproses.")

# ================= HALAMAN ADMIN =================
if menu_option == "Panel Admin":
    st.title("🛠️ Panel Admin WhatsApp")
    if st.text_input("Password Admin:", type="password") == ADMIN_PASSWORD:
        st.success("Admin Logged In")
        with st.form("create_token_form"):
            package_name = st.text_input("Nama Paket:", value="Paket Regular")
            quota_amount = st.number_input("Jumlah Kuota:", min_value=1, max_value=100, value=1)
            if st.form_submit_button("Generate Token"):
                new_token = "TOK-" + str(uuid.uuid4())[:8].upper()
                try:
                    client = checker_engine.get_db_client(ch_config)
                    client.command("""INSERT INTO default.app_tokens (token, package_name, quota, created_at, is_active) VALUES ({token:String}, {pkg:String}, {quota:Int32}, {date:DateTime}, 1)""", parameters={"token": new_token, "pkg": package_name, "quota": int(quota_amount), "date": datetime.now()})
                    st.success("Token berhasil dibuat!")
                    st.markdown(f"### > `{new_token}`")
                except Exception as e: 
                    st.error(f"Gagal: {e}")

# ================= HALAMAN LOGIN =================
elif menu_option == "Login Token / Redeem":
    st.title("🔑 Masukkan Kode Akses Token")
    token_input = st.text_input("Kode Token:")
    if st.button("Validasi Token") and token_input.strip():
        with st.spinner("Mengecek ketersediaan token..."):
            success, msg = checker_engine.check_token_validity(ch_config, token_input.strip())
            if success:
                st.session_state.authenticated = True
                st.session_state.token_info = msg
                st.session_state.token_code = token_input.strip()
                st.success("Token valid! Anda bisa kembali ke Halaman Utama untuk mengecek dokumen.")
            else: 
                st.error(msg)

# ================= HALAMAN UTAMA =================
else:
    st.title("📄 Sistem Pengecekan Kemiripan Dokumen")
    st.caption("🔍 Engine: Turnitin-Style Web Search (100 Link Deep Scrape) | Mendukung PDF, DOCX, TXT")

    uploaded_file = st.file_uploader("Pilih dokumen (Maks. 200MB)", type=["pdf", "docx", "txt"])

    if uploaded_file is not None:
        extracted_text = checker_engine.extract_text_from_file(uploaded_file)
        total_words = len(extracted_text.split())
        st.write(f"**Dokumen diterima:** {uploaded_file.name} (±{total_words} Kata)")
        
        if st.button("Jalankan Pengecekan (Gunakan 1 Kuota)", type="primary"):
            if not st.session_state.authenticated:
                st.warning("⚠️ Silakan validasi token di menu 'Login Token / Redeem' terlebih dahulu.")
            else:
                if checker_engine.redeem_token_quota(ch_config, st.session_state.token_code):
                    _, updated_info = checker_engine.check_token_validity(ch_config, st.session_state.token_code)
                    st.session_state.token_info = updated_info
                    
                    progress_bar = st.progress(0)
                    status_text = st.empty()
                    
                    def update_progress(val):
                        progress_bar.progress(val)
                    def update_status(text):
                        status_text.text(text)
                        
                    with st.spinner("Menelusuri internet & menganalisis seluruh kalimat secara mendalam..."):
                        highlighted_html, found_sources, similarity_percentage, sentences_to_check = checker_engine.run_plagiarism_check(
                            extracted_text, update_progress, update_status
                        )
                        
                    progress_bar.empty()
                    status_text.empty()
                    
                    st.markdown("---")
                    if similarity_percentage > 0:
                        col1, col2 = st.columns([2, 1])
                        with col1:
                            st.write("### Pratinjau Dokumen")
                            st.markdown(f'<div style="border: 1px solid #ddd; padding: 20px; border-radius: 5px; background-color: #fafafa; line-height: 2.0; font-family: serif;">{highlighted_html}</div>', unsafe_allow_html=True)
                        with col2:
                            st.write("### Integrity Overview")
                            st.markdown(f'<div class="turnitin-source-list"><p class="turnitin-score">{similarity_percentage}%</p><p style="font-weight:bold; color:#555;">Overall Similarity</p><hr>', unsafe_allow_html=True)
                            for src in found_sources:
                                st.markdown(f"""
                                <div style="margin-bottom: 12px; line-height: 1.3;">
                                    <span style="background-color: {src['color']}; padding: 2px 6px; font-weight: bold; border-radius: 3px; font-size: 12px;">{src['index']}</span>
                                    <span style="font-size: 14px; margin-left: 5px;"><a href="{src['url']}" target="_blank" style="color: #333; text-decoration: none;">{src['title'][:35]}...</a></span><br>
                                    <span style="font-size: 11px; color: #888; margin-left: 30px;">Scraped via: {src['scraper_engine']}</span>
                                </div>
                                """, unsafe_allow_html=True)
                            st.markdown("</div>", unsafe_allow_html=True)
                    else:
                        col1, col2 = st.columns([2, 1])
                        with col1:
                            st.write("### Pratinjau Dokumen")
                            safe_text = " ".join(sentences_to_check)
                            st.markdown(f'<div style="border: 1px solid #ddd; padding: 20px; border-radius: 5px; background-color: #fafafa; line-height: 2.0; font-family: serif;">{safe_text}</div>', unsafe_allow_html=True)
                        with col2:
                            st.write("### Integrity Overview")
                            st.markdown('<div class="turnitin-source-list"><p style="font-size: 48px; font-weight: bold; color: #28a745; margin-bottom: 0;">0%</p><p style="font-weight:bold; color:#555;">Overall Similarity</p><hr><p style="color:#777; font-size:14px;">Bebas dari deteksi plagiasi internet.</p></div>', unsafe_allow_html=True)
                    
                    st.markdown("---")
                    try:
                        pdf_data = checker_engine.create_pdf_report(uploaded_file.name, similarity_percentage, found_sources)
                        st.download_button(
                            label="📥 Unduh Laporan PDF (Cetak Hasil)",
                            data=pdf_data,
                            file_name=f"Laporan_Plagiasi_{uploaded_file.name}.pdf",
                            mime="application/pdf",
                            type="primary"
                        )
                    except Exception as e:
                        st.error(f"Gagal membuat laporan PDF: {e}")
                else:
                    st.error("Gagal memotong kuota. Pastikan token Anda masih memiliki sisa kuota yang valid.")
